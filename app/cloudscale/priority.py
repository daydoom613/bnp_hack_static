"""Request classification from data/service_priority.xlsx (sheet "Priority")."""
import logging
from pathlib import Path

from . import config

log = logging.getLogger(__name__)

DEFAULTS = {
    "CreateOrder": {"priority": 1, "critical": True},
    "UpdateItem": {"priority": 2, "critical": True},
    "GetCatalog": {"priority": 3, "critical": False},
}


def _truthy(value):
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "yes"}
    return bool(value)


def load(path=None):
    path = Path(path or config.PRIORITY_FILE)
    if not path.exists():
        log.warning("%s not found, using built-in priorities", path)
        return dict(DEFAULTS)

    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True)["Priority"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    col = {name: header.index(name) for name in ("request_type", "priority", "critical")}

    table = {}
    for row in rows:
        if row[col["request_type"]] is None:
            continue
        table[str(row[col["request_type"]]).strip()] = {
            "priority": int(row[col["priority"]]),
            "critical": _truthy(row[col["critical"]]),
        }
    log.info("loaded %d request types from %s", len(table), path)
    return table


TABLE = load()


def is_critical(request_type):
    """Unknown request types count as critical: the safe side of the Spot rule."""
    entry = TABLE.get(request_type)
    return True if entry is None else entry["critical"]


def priority_of(request_type):
    entry = TABLE.get(request_type)
    return 1 if entry is None else entry["priority"]
