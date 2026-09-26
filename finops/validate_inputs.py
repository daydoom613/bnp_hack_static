#!/usr/bin/env python3
"""Validate the challenge's input files against the Dataset Building Document.

    python finops/validate_inputs.py            # exit 1 if any file is invalid

  traffic_profiles      UTF-8, comma-separated, header exactly
                        time_offset_min,rps_target,duration_sec,concurrency; integer rows
  budget_cap.txt        one numeric line; lines starting with # are ignored
  service_priority      .xlsx with a sheet named "Priority" and the case-sensitive columns
                        request_type (text), priority (integer, 1 = highest), critical (TRUE/FALSE)
  pricing_matrix        region,instance_type,on_demand_price_per_hour,reserved_price_per_hour,
                        spot_price_per_hour; includes every instance type Terraform launches
  security_baseline     .docx naming the four mandatory tags
  baseline_cost.json    valid UTF-8 JSON with a numeric TotalMonthlyCost

Every traffic profile in data/ and finops-app/data/ is checked.
"""
import csv
import json
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
PROFILE_HEADER = ["time_offset_min", "rps_target", "duration_sec", "concurrency"]
PRICING_HEADER = ["region", "instance_type", "on_demand_price_per_hour", "reserved_price_per_hour",
                  "spot_price_per_hour"]
MANDATORY_TAGS = ["Owner", "Project", "Environment", "Cost_Center"]

errors = []


def check(ok, message):
    if not ok:
        errors.append(message)
    return ok


def read_utf8(path):
    try:
        return path.read_bytes().decode("utf-8")
    except UnicodeDecodeError:
        errors.append(f"{path.name}: not UTF-8")
        return ""


def traffic_profile(path):
    rows = list(csv.reader(read_utf8(path).splitlines()))
    if not check(rows and [c.strip() for c in rows[0]] == PROFILE_HEADER,
                 f"{path.name}: header must be exactly {','.join(PROFILE_HEADER)}"):
        return
    check(len(rows) > 1, f"{path.name}: no slices")
    for n, row in enumerate(rows[1:], start=2):
        if not row:
            continue
        if not check(len(row) == 4 and all(re.fullmatch(r"\d+", c.strip()) for c in row),
                     f"{path.name} line {n}: four non-negative integers expected, got {row}"):
            continue
        _, rps, duration, concurrency = (int(c) for c in row)
        check(rps > 0 and duration > 0 and concurrency > 0,
              f"{path.name} line {n}: rps_target, duration_sec and concurrency must be > 0")


def budget_cap(path):
    lines = [l.strip() for l in read_utf8(path).splitlines() if l.strip() and not l.strip().startswith("#")]
    check(len(lines) == 1 and re.fullmatch(r"\d+(\.\d+)?", lines[0] if lines else ""),
          f"{path.name}: exactly one numeric line expected, got {lines}")


def service_priority(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if not check("Priority" in workbook.sheetnames, f"{path.name}: no sheet named 'Priority' ({workbook.sheetnames})"):
        return
    rows = list(workbook["Priority"].iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]] if rows else []
    missing = [c for c in ("request_type", "priority", "critical") if c not in header]
    if not check(not missing, f"{path.name}: Priority sheet is missing columns {missing}"):
        return
    col = {name: header.index(name) for name in ("request_type", "priority", "critical")}
    data = [r for r in rows[1:] if r and r[col["request_type"]] is not None]
    check(data, f"{path.name}: no request types")
    for r in data:
        name = r[col["request_type"]]
        check(isinstance(r[col["priority"]], int) and r[col["priority"]] >= 1,
              f"{path.name}: {name}: priority must be an integer >= 1")
        check(isinstance(r[col["critical"]], bool) or str(r[col["critical"]]).upper() in ("TRUE", "FALSE"),
              f"{path.name}: {name}: critical must be TRUE or FALSE")


def terraform_instance_types():
    """Instance types the Terraform environments launch (tfvars values and variable defaults)."""
    found = set()
    for path in (ROOT / "terraform" / "environments").glob("*/*.tf*"):
        if path.suffix not in (".tf", ".tfvars"):
            continue
        text = path.read_text(encoding="utf-8")
        for match in re.finditer(r'instance_type\w*\s*=\s*"([a-z0-9]+\.[a-z0-9]+)"', text):
            found.add(match.group(1))
        for block in re.finditer(r'variable\s+"\w*instance_type\w*"\s*{[^}]*?default\s*=\s*"([a-z0-9]+\.[a-z0-9]+)"', text, re.S):
            found.add(block.group(1))
    return found


def pricing_matrix(path):
    rows = list(csv.reader(read_utf8(path).splitlines()))
    if not check(rows and [c.strip() for c in rows[0]] == PRICING_HEADER,
                 f"{path.name}: header must be exactly {','.join(PRICING_HEADER)}"):
        return
    priced = set()
    for n, row in enumerate(rows[1:], start=2):
        if not row:
            continue
        try:
            od, ri, spot = (float(x) for x in row[2:5])
        except ValueError:
            errors.append(f"{path.name} line {n}: prices must be numbers")
            continue
        check(0 < spot <= od and 0 < ri <= od, f"{path.name} line {n}: expected 0 < spot, reserved <= on-demand")
        priced.add(row[1].strip())
    missing = sorted(terraform_instance_types() - priced)
    check(not missing, f"{path.name}: instance types used by Terraform are not priced: {missing}")


def security_baseline(path):
    with zipfile.ZipFile(path) as docx:
        text = re.sub(r"<[^>]+>", "", docx.read("word/document.xml").decode("utf-8"))
    missing = [t for t in MANDATORY_TAGS if t not in text]
    check(not missing, f"{path.name}: mandatory tags not described: {missing}")


def baseline_cost(path):
    try:
        doc = json.loads(read_utf8(path))
    except ValueError as exc:
        errors.append(f"{path.name}: invalid JSON ({exc})")
        return
    check(isinstance(doc.get("TotalMonthlyCost"), (int, float)), f"{path.name}: numeric TotalMonthlyCost missing")


def main():
    checks = [
        (DATA / "budget_cap.txt", budget_cap),
        (DATA / "service_priority.xlsx", service_priority),
        (DATA / "pricing_matrix.csv", pricing_matrix),
        (DATA / "security_baseline.docx", security_baseline),
        (DATA / "baseline_cost.json", baseline_cost),
    ]
    checks += [(p, service_priority) for p in sorted(DATA.glob("service_priority_*.xlsx"))]
    checks += [(p, traffic_profile) for p in sorted(DATA.glob("traffic_profile*.csv"))]
    checks += [(p, traffic_profile) for p in sorted((ROOT / "finops-app" / "data").glob("traffic_profile*.csv"))]

    for path, validate in checks:
        before = len(errors)
        if check(path.exists(), f"{path.relative_to(ROOT)}: missing"):
            validate(path)
        print(f"{'ok  ' if len(errors) == before else 'FAIL'}  {path.relative_to(ROOT).as_posix()}")

    for message in errors:
        print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
